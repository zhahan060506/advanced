from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path(r"F:\ESP\shoubiaob\第四章_智能小电视系统集成与测试.docx")


def set_run_font(run, size=None, bold=None):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def format_paragraph(paragraph, first_indent=True):
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if first_indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size={1: 16, 2: 14, 3: 12}.get(level, 12), bold=True)
    return paragraph


def add_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=12)
    return paragraph


def add_table(doc, title, headers, rows):
    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(title_paragraph, first_indent=False)
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, size=10.5)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)

    doc.add_paragraph("")
    return table


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, bold=bold)


def build_doc():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.color.rgb = None

    add_heading(doc, "第四章  智能小电视系统集成与测试", 1)

    add_heading(doc, "4.1 模块/单元/子系统测试", 2)
    add_paragraph(
        doc,
        "为了验证本项目各功能模块能否稳定运行，本章分别对显示模块、音频模块、语音交互模块、倒计时模块和外设控制模块进行了测试。"
        "测试过程中，设备通过 ESP-IDF 完成程序烧录，并使用串口监视器观察系统运行日志，结合实际硬件现象判断功能是否正常。"
    )

    add_heading(doc, "4.1.1 显示模块测试", 3)
    add_paragraph(
        doc,
        "显示模块采用 1.54 寸 ST7789 TFT 彩色屏，使用 SPI 接口与 ESP32-S3 通信。测试时，首先观察设备上电后的屏幕初始化过程，"
        "确认屏幕是否能够正常点亮，随后检查主页、聊天页和倒计时页是否能够正常切换。"
    )
    add_paragraph(
        doc,
        "测试结果表明，设备上电后屏幕能够正常点亮并进入初始化界面。系统连接网络并完成服务器登录后，屏幕自动切换到主页。"
        "主页能够显示当前时间、日期、星期、天气、电量和网络状态。进入聊天状态后，屏幕能够切换到聊天页面；"
        "启动倒计时后，屏幕能够切换到倒计时页面并显示剩余时间。"
    )
    add_table(doc, "表4-1 显示模块测试结果", ["测试项目", "测试方法", "测试结果"], [
        ["屏幕点亮", "设备上电观察背光和画面", "正常点亮"],
        ["主页显示", "登录服务器后观察主页内容", "时间、日期、天气等显示正常"],
        ["聊天页切换", "唤醒词唤醒后观察页面", "能够切换到聊天页"],
        ["倒计时页显示", "语音设置倒计时", "能够显示倒计时页面"],
        ["页面返回", "AI 回答结束或倒计时结束", "能够返回主页"],
    ])

    add_heading(doc, "4.1.2 音频模块测试", 3)
    add_paragraph(
        doc,
        "音频模块包括麦克风输入和扬声器输出两部分。麦克风采用 ICS-43434 数字 I2S 麦克风，扬声器部分采用 MAX98357A I2S 功放模块。"
        "测试时，通过串口日志观察 I2S 初始化状态，并通过实际语音输入和语音播放验证音频链路是否正常。"
    )
    add_paragraph(
        doc,
        "在早期测试中，麦克风存在无法正常采集数据的问题，表现为语音识别结果异常。经过排查后，发现主要原因是 I2S 工作模式和麦克风引脚配置不匹配。"
        "项目最终将音频配置调整为 I2S Simplex 模式，使麦克风和扬声器分别使用独立的 I2S 通道。修改后，麦克风能够正常采集语音，"
        "扬声器能够播放系统提示音和 AI 语音回复。"
    )
    add_table(doc, "表4-2 音频模块测试结果", ["测试项目", "测试方法", "测试结果"], [
        ["扬声器播放", "启动提示音、AI 回答播放", "播放正常"],
        ["麦克风采集", "说话后观察识别结果", "能够正常识别"],
        ["I2S 时钟信号", "检查麦克风 WS/SCK 工作状态", "修改后正常"],
        ["音量调节", "使用 KEY2 调节音量", "音量变化正常"],
    ])

    add_heading(doc, "4.1.3 唤醒词与语音交互测试", 3)
    add_paragraph(
        doc,
        "本项目将原本按键进入聊天的方式改为唤醒词进入聊天。设备在主页状态下不会直接进入语音聊天，而是只监听唤醒词。"
        "当用户说出“你好小智”后，系统进入聊天页面并开启语音输入通道。"
    )
    add_paragraph(
        doc,
        "测试时，首先在主页状态下直接说普通语句，观察设备是否误进入聊天。随后说出唤醒词，观察屏幕是否切换到聊天页，"
        "并通过串口日志确认系统状态是否变为 listening。测试结果表明，设备在主页状态下不会直接响应普通语音，只有检测到唤醒词后才进入聊天模式。"
        "AI 回答完成后，系统等待约 5 秒自动返回主页，并重新进入待机状态。"
    )
    add_table(doc, "表4-3 语音交互测试结果", ["测试项目", "测试方法", "测试结果"], [
        ["主页普通语音", "在主页直接说问题", "不进入聊天"],
        ["唤醒词检测", "说“你好小智”", "成功进入聊天页"],
        ["语音问答", "唤醒后提出问题", "AI 能够回复"],
        ["自动返回主页", "AI 回答结束后等待", "约 5 秒后回主页"],
        ["主页待机状态", "回主页后再次普通说话", "不会继续聊天"],
    ])

    add_heading(doc, "4.1.4 倒计时模块测试", 3)
    add_paragraph(
        doc,
        "倒计时模块是本项目新增功能之一。用户可以通过语音说出“设置一个 60 秒倒计时”等指令，系统识别后调用本地倒计时工具，"
        "屏幕切换到倒计时页面，并实时显示剩余秒数。"
    )
    add_paragraph(
        doc,
        "测试时，分别设置 3 秒、10 秒和 60 秒倒计时，观察倒计时页面是否正常显示，时间是否逐秒递减，到时后是否能够响铃提示。"
        "测试结果表明，倒计时功能能够正常启动，页面能够正确跳转，倒计时结束后扬声器能够播放提示音。按下 KEY1 时，可以取消倒计时或停止响铃。"
    )
    add_table(doc, "表4-4 倒计时模块测试结果", ["测试项目", "测试方法", "测试结果"], [
        ["语音设置倒计时", "说“设置一个 3 秒倒计时”", "能够启动"],
        ["页面跳转", "倒计时启动后观察屏幕", "能切换到倒计时页"],
        ["时间递减", "观察倒计时数字", "能够逐秒变化"],
        ["到时响铃", "倒计时结束后观察扬声器", "能够响铃"],
        ["按键取消", "倒计时中按 KEY1", "能够取消并返回主页"],
    ])

    add_heading(doc, "4.1.5 LED 控制模块测试", 3)
    add_paragraph(
        doc,
        "本项目增加了语音控制 LED 小灯功能。LED 接在 ESP32-S3 的 IO10 引脚，软件中通过 GPIO 输出高低电平控制 LED 的亮灭。"
        "测试时，用户通过语音说出“开灯”和“关灯”，观察 LED 状态是否发生变化。"
    )
    add_paragraph(
        doc,
        "测试结果表明，系统能够根据 AI 识别出的语音指令调用本地灯控工具。当用户说“开灯”时，IO10 输出高电平，LED 点亮；"
        "当用户说“关灯”时，IO10 输出低电平，LED 熄灭。"
    )
    add_table(doc, "表4-5 LED 控制测试结果", ["测试项目", "测试方法", "测试结果"], [
        ["开灯", "语音说“开灯”", "LED 点亮"],
        ["关灯", "语音说“关灯”", "LED 熄灭"],
        ["GPIO 输出", "测试 IO10 电平变化", "电平变化正常"],
    ])

    add_heading(doc, "4.2 系统集成测试", 2)
    add_paragraph(
        doc,
        "在完成各模块单独测试后，对系统进行了整体集成测试。集成测试主要验证显示、语音、网络通信、倒计时、按键和 LED 控制等功能是否能够协同工作。"
    )
    add_paragraph(
        doc,
        "系统集成测试流程如下：设备上电后完成屏幕、音频、按键和网络初始化；随后系统连接 WiFi，并连接小智 AI 云端服务器；"
        "连接成功后屏幕进入主页，显示时间、天气、日期和电量信息；用户说出唤醒词后，设备进入聊天页面；用户进行语音提问，"
        "设备完成语音识别、云端交互和语音播放；AI 回答结束后，设备自动返回主页；用户通过语音设置倒计时，设备切换倒计时页面并到时响铃；"
        "用户通过语音控制 LED，实现开灯和关灯。"
    )
    add_paragraph(
        doc,
        "测试结果表明，各模块之间能够正常配合工作。系统在主页状态下保持待机，只响应唤醒词；进入聊天状态后能够进行语音交互；"
        "倒计时和 LED 控制功能能够通过语音触发；AI 回答结束或倒计时结束后，系统能够返回主页。"
    )
    add_table(doc, "表4-6 系统集成测试结果", ["测试内容", "预期结果", "实际结果"], [
        ["系统启动", "正常初始化并进入主页", "符合预期"],
        ["网络连接", "成功连接服务器", "符合预期"],
        ["唤醒词进入聊天", "说唤醒词后进入聊天页", "符合预期"],
        ["AI 语音问答", "能识别语音并播放回答", "符合预期"],
        ["回主页逻辑", "回主页后不继续聊天", "符合预期"],
        ["倒计时功能", "能设置、显示、响铃", "符合预期"],
        ["LED 控制", "能语音开关灯", "符合预期"],
    ])

    add_heading(doc, "4.3 实物展示", 2)
    add_paragraph(
        doc,
        "本项目最终完成的智能小电视实物由 ESP32-S3 主控板、1.54 寸 ST7789 彩色屏、ICS-43434 麦克风、MAX98357A 功放模块、"
        "扬声器、按键和 LED 小灯组成。设备能够完成主页信息显示、语音唤醒、AI 对话、倒计时提醒和语音控制 LED 等功能。"
    )
    add_paragraph(
        doc,
        "实物运行时，主页显示当前时间、日期、星期、天气和电量等信息，如图 4-1 所示。用户说出唤醒词后，设备进入聊天页面，"
        "能够进行语音问答，如图 4-2 所示。语音设置倒计时后，屏幕切换到倒计时页面，并在倒计时结束后播放提示音，如图 4-3 所示。"
        "语音控制 LED 时，外接 LED 能够根据“开灯”“关灯”指令进行亮灭变化，如图 4-4 所示。"
    )
    add_paragraph(
        doc,
        "通过实物测试可以看出，本项目完成了基于小智 AI 官方程序的二次开发，实现了适配本组硬件平台的智能语音交互设备。"
        "系统整体运行稳定，能够满足课程项目中对语音交互、显示反馈和外设控制的设计要求。"
    )

    return doc


if __name__ == "__main__":
    document = build_doc()
    document.save(OUT)
    print(OUT)
